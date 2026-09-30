"""Impor data HariLibur dari berkas Excel (.xlsx) atau CSV.

Dipakai bersama oleh menu Setting Kalender (Admin Biro PAKLN, unggah
Excel) dan management command `seed_hari_libur` (CSV bawaan), supaya
aturan validasinya sama.

Format berkas: baris header berisi kolom `tanggal`, `keterangan`, `jenis`
(urutan bebas, huruf besar/kecil diabaikan), lalu satu baris per tanggal.
"""

import csv
import io
from datetime import date, datetime
from typing import NamedTuple

from django.db import transaction

from .models import HariLibur

KOLOM = ("tanggal", "keterangan", "jenis")
FORMAT_TANGGAL_TEKS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y")
MAKS_BARIS = 1000

# Variasi penulisan jenis yang diterima -> nilai HariLibur.Jenis.
_ALIAS_JENIS = {
    "libur_nasional": HariLibur.Jenis.LIBUR_NASIONAL,
    "libur nasional": HariLibur.Jenis.LIBUR_NASIONAL,
    "libur": HariLibur.Jenis.LIBUR_NASIONAL,
    "cuti_bersama": HariLibur.Jenis.CUTI_BERSAMA,
    "cuti bersama": HariLibur.Jenis.CUTI_BERSAMA,
    "cuti": HariLibur.Jenis.CUTI_BERSAMA,
}


class BarisLibur(NamedTuple):
    tanggal: date
    keterangan: str
    jenis: str


class ImporError(Exception):
    """Berkas tidak valid; `errors` berisi pesan per baris untuk admin."""

    def __init__(self, errors):
        self.errors = errors
        super().__init__("; ".join(errors))


def _parse_tanggal(nilai):
    if isinstance(nilai, datetime):
        return nilai.date()
    if isinstance(nilai, date):
        return nilai
    teks = str(nilai or "").strip()
    for fmt in FORMAT_TANGGAL_TEKS:
        try:
            return datetime.strptime(teks, fmt).date()
        except ValueError:
            continue
    return None


def _parse_baris(rows, sumber):
    """`rows` = iterable tuple sel (baris pertama yang tidak kosong adalah
    header). Mengumpulkan SEMUA kesalahan dulu, baru gagal sekaligus —
    admin bisa memperbaiki berkasnya dalam sekali jalan."""
    rows = iter(rows)
    header = None
    no_header = 0
    for no_header, row in enumerate(rows, start=1):
        if any(str(c or "").strip() for c in row):
            header = [str(c or "").strip().lower() for c in row]
            break
    if header is None:
        raise ImporError([f"{sumber}: berkas kosong."])
    hilang = [k for k in KOLOM if k not in header]
    if hilang:
        raise ImporError([f"{sumber}: kolom {', '.join(hilang)} tidak ditemukan pada baris header."])
    idx = {k: header.index(k) for k in KOLOM}

    hasil, errors, dilihat = [], [], {}
    for no, row in enumerate(rows, start=no_header + 1):
        row = list(row) + [None] * (len(header) - len(row))
        if not any(str(c or "").strip() for c in row):
            continue
        if len(hasil) + len(errors) >= MAKS_BARIS:
            errors.append(f"{sumber}: maksimal {MAKS_BARIS} baris per berkas.")
            break

        mentah_tanggal = row[idx["tanggal"]]
        tanggal = _parse_tanggal(mentah_tanggal)
        keterangan = str(row[idx["keterangan"]] or "").strip()
        jenis_teks = str(row[idx["jenis"]] or "").strip().lower()
        jenis = _ALIAS_JENIS.get(jenis_teks) if jenis_teks else HariLibur.Jenis.LIBUR_NASIONAL

        salah = []
        if tanggal is None:
            salah.append(f"tanggal {mentah_tanggal!r} tidak valid (gunakan YYYY-MM-DD atau DD-MM-YYYY)")
        elif tanggal in dilihat:
            salah.append(f"tanggal {tanggal:%d-%m-%Y} ganda (juga di baris {dilihat[tanggal]})")
        if not keterangan:
            salah.append("keterangan kosong")
        elif len(keterangan) > 150:
            salah.append("keterangan lebih dari 150 karakter")
        if jenis is None:
            salah.append(f"jenis {jenis_teks!r} tidak dikenal (Libur Nasional / Cuti Bersama)")

        if salah:
            errors.append(f"{sumber} baris {no}: " + "; ".join(salah))
            continue
        dilihat[tanggal] = no
        hasil.append(BarisLibur(tanggal, keterangan, jenis))

    if errors:
        raise ImporError(errors)
    return hasil


def baca_xlsx(file_obj, sumber="Excel"):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from zipfile import BadZipFile

    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError):
        raise ImporError([f"{sumber}: berkas bukan Excel (.xlsx) yang valid."])
    try:
        return _parse_baris(wb.worksheets[0].iter_rows(values_only=True), sumber)
    finally:
        wb.close()


def baca_csv(file_obj, sumber="CSV"):
    """`file_obj` boleh teks atau biner. Baris yang diawali '#' diabaikan."""
    konten = file_obj.read()
    if isinstance(konten, bytes):
        try:
            konten = konten.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ImporError([f"{sumber}: berkas CSV harus berenkode UTF-8."])
    # Komentar diganti baris kosong (bukan dibuang) agar nomor baris pada
    # pesan kesalahan tetap sesuai berkas aslinya.
    lines = ["\n" if line.lstrip().startswith("#") else line for line in io.StringIO(konten)]
    return _parse_baris(csv.reader(lines), sumber)


def baca_berkas(uploaded):
    """Berkas unggahan Django -> list BarisLibur, berdasarkan ekstensi."""
    nama = uploaded.name.lower()
    if nama.endswith(".xlsx"):
        return baca_xlsx(uploaded, uploaded.name)
    if nama.endswith(".csv"):
        return baca_csv(uploaded, uploaded.name)
    raise ImporError([f"{uploaded.name}: format tidak didukung (gunakan .xlsx atau .csv)."])


def simpan(baris, timpa=False, user=None):
    """Simpan hasil baca ke DB dalam satu transaksi. Tanggal yang sudah ada
    dilewati, atau — bila `timpa` — keterangan & jenisnya diperbarui
    (status aktif tidak diubah)."""
    ada = {h.tanggal: h for h in HariLibur.objects.filter(tanggal__in=[b.tanggal for b in baris])}
    baru, diperbarui, dilewati = [], [], []
    for b in baris:
        lama = ada.get(b.tanggal)
        if lama is None:
            baru.append(HariLibur(tanggal=b.tanggal, keterangan=b.keterangan, jenis=b.jenis, dibuat_oleh=user))
        elif timpa and (lama.keterangan, lama.jenis) != (b.keterangan, b.jenis):
            lama.keterangan, lama.jenis = b.keterangan, b.jenis
            diperbarui.append(lama)
        else:
            dilewati.append(b.tanggal)

    with transaction.atomic():
        HariLibur.objects.bulk_create(baru)
        for h in diperbarui:
            h.save(update_fields=["keterangan", "jenis", "updated_at"])

    return {"dibuat": len(baru), "diperbarui": len(diperbarui), "dilewati": len(dilewati)}


def buat_template_xlsx():
    """Template Excel kosong (dengan contoh & dropdown jenis) untuk diunduh admin."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Hari Libur"
    ws.append(list(KOLOM))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    ws.append([date(2026, 8, 17), "Proklamasi Kemerdekaan", "Libur Nasional"])
    ws.append([date(2026, 12, 24), "Cuti Bersama Kelahiran Yesus Kristus", "Cuti Bersama"])
    for row in ws.iter_rows(min_row=2, max_col=1):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 18

    dv = DataValidation(type="list", formula1='"Libur Nasional,Cuti Bersama"', allow_blank=True)
    dv.error = "Pilih Libur Nasional atau Cuti Bersama."
    ws.add_data_validation(dv)
    dv.add("C2:C1000")

    petunjuk = wb.create_sheet("Petunjuk")
    for baris in [
        ["Petunjuk pengisian"],
        ["tanggal", "Tanggal libur. Format sel tanggal Excel, atau teks YYYY-MM-DD / DD-MM-YYYY."],
        ["keterangan", "Nama hari libur, maks. 150 karakter."],
        ["jenis", "Libur Nasional atau Cuti Bersama (kosong = Libur Nasional)."],
        [],
        ["Satu baris = satu tanggal. Sabtu & Minggu tidak perlu diisi."],
        ["Hapus baris contoh sebelum mengunggah jika tidak diperlukan."],
        ["Tanggal yang sudah ada dilewati, kecuali opsi 'Timpa data yang sudah ada' dicentang saat impor."],
    ]:
        petunjuk.append(baris)
    petunjuk["A1"].font = Font(bold=True)
    petunjuk.column_dimensions["A"].width = 14
    petunjuk.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
