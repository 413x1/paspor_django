"""
Report terpadu Dinas & Non-Dinas — wiki/instructions/BISNIS_PROSES_PDLN.MD §10.

- Cakupan per role diterapkan di queryset (`queryset_untuk`), bukan di template:
    Admin Unor  : unit organisasinya, semua jenis
    Admin BPSDM : PDLN Tipe 2 (T2P/T2L), semua unor
    Admin PAKLN : semua
- Filter global (`terapkan_filter`) dipakai Dasbor, Export, dan Rekap.
"""

import csv
import io
from collections import Counter, defaultdict
from datetime import date
from statistics import median

from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.utils import timezone

from accounts.models import User

from .models import LaporanPdln, Pengajuan, PermohonanPembatalan, RiwayatPengajuan

S = Pengajuan.Status
T = Pengajuan.TipePdln
JP = Pengajuan.JenisPerjalanan
Aksi = RiwayatPengajuan.Aksi

BATAS_HARI_LAPORAN = 14  # §14 pertanyaan terbuka #4 (default, konfigurabel di sini)


def q_unor(unit_organisasi_id):
    """Pengajuan milik unit organisasi — snapshot di pengajuan, atau profil
    pegawai bila snapshot belum terisi (draft lama)."""
    return Q(unit_organisasi_id=unit_organisasi_id) | Q(
        unit_organisasi__isnull=True, pegawai__profile__unit_organisasi_id=unit_organisasi_id,
    )


def queryset_untuk(user, termasuk_draft=False):
    qs = Pengajuan.objects.all()
    if not termasuk_draft:
        # Pengajuan yang pernah dikirim (termasuk yang sedang dikembalikan
        # ke pegawai) atau yang sudah dibatalkan setelah dikirim.
        qs = qs.exclude(tgl_pengajuan__isnull=True)
    if user.role == User.Role.ADMIN_UNOR:
        if not user.unit_organisasi_id:
            return qs.none()
        qs = qs.filter(q_unor(user.unit_organisasi_id))
    elif user.role == User.Role.ADMIN_BPSDM:
        qs = qs.filter(jenis_perjalanan=JP.PDLN, tipe_pdln__in=[T.T2P, T.T2L])
    elif user.role != User.Role.ADMIN_PAKLN:
        return qs.none()
    return qs


def _tanggal(nilai):
    try:
        return date.fromisoformat(nilai) if nilai else None
    except ValueError:
        return None


def baca_filter(params):
    """Normalisasi parameter GET filter global."""
    return {
        "jenis": params.get("jenis", "") if params.get("jenis", "") in JP.values else "",
        "tipe": [t for t in params.getlist("tipe") if t in T.values] if hasattr(params, "getlist") else [],
        "status": [s for s in params.getlist("status") if s in S.values] if hasattr(params, "getlist") else [],
        "unor": params.get("unor", ""),
        "dari": _tanggal(params.get("dari")),
        "sampai": _tanggal(params.get("sampai")),
        "basis_tanggal": "tgl_berangkat" if params.get("basis_tanggal") == "tgl_berangkat" else "tgl_pengajuan",
        "negara": params.get("negara", ""),
        "kategori": params.get("kategori", ""),
        "sumber": params.get("sumber", ""),
        "kanal": params.get("kanal", "") if params.get("kanal", "") in Pengajuan.Kanal.values else "",
    }


def terapkan_filter(qs, f):
    if f["jenis"]:
        qs = qs.filter(jenis_perjalanan=f["jenis"])
    if f["tipe"]:
        qs = qs.filter(tipe_pdln__in=f["tipe"])
    if f["status"]:
        qs = qs.filter(status__in=f["status"])
    if f["unor"].isdigit():
        qs = qs.filter(q_unor(int(f["unor"])))
    if f["dari"]:
        qs = qs.filter(**{f"{f['basis_tanggal']}__gte": f["dari"]})
    if f["sampai"]:
        qs = qs.filter(**{f"{f['basis_tanggal']}__lte": f["sampai"]})
    if f["negara"].isdigit():
        qs = qs.filter(tujuan_negara__id=int(f["negara"])).distinct()
    if f["kategori"].isdigit():
        qs = qs.filter(kategori_id=int(f["kategori"]))
    if f["sumber"].isdigit():
        qs = qs.filter(sumber_pembiayaan_id=int(f["sumber"]))
    if f["kanal"]:
        qs = qs.filter(kanal=f["kanal"])
    return qs


# ---------------------------------------------------------------------------
# Dasbor
# ---------------------------------------------------------------------------

STATUS_ROLE = {
    User.Role.ADMIN_UNOR: S.PROSES,
    User.Role.ADMIN_BPSDM: S.PROSES_BPSDM,
    User.Role.ADMIN_PAKLN: S.PROSES_PAKLN,
}


def pembatalan_menunggu(user, qs_pengajuan=None):
    """Permohonan pembatalan yang menunggu keputusan role `user`."""
    qs = PermohonanPembatalan.objects.filter(pengajuan__in=qs_pengajuan if qs_pengajuan is not None else queryset_untuk(user))
    if user.role == User.Role.ADMIN_UNOR:
        return qs.filter(status=PermohonanPembatalan.Status.MENUNGGU_UNOR)
    if user.role == User.Role.ADMIN_PAKLN:
        return qs.filter(status=PermohonanPembatalan.Status.MENUNGGU_PAKLN)
    return qs.none()


def ringkasan(user, qs):
    hitung = dict(qs.values_list("status").annotate(n=Count("id")).values_list("status", "n"))
    status_saya = STATUS_ROLE.get(user.role)
    perlu = qs.filter(status=status_saya).exclude(
        permohonan_pembatalan__status__in=PermohonanPembatalan.STATUS_TERBUKA
    ).count() if status_saya else 0
    batal_menunggu = pembatalan_menunggu(user, qs).count()
    laporan_menunggu = 0
    if user.role == User.Role.ADMIN_PAKLN:
        laporan_menunggu = LaporanPdln.objects.filter(
            pengajuan__in=qs, status=LaporanPdln.Status.MENUNGGU
        ).count()
    dikembalikan = qs.filter(status=S.BELUM).count()
    return {
        "total": qs.count(),
        "perlu_tindakan": perlu + batal_menunggu + laporan_menunggu,
        "pembatalan_menunggu": batal_menunggu,
        "laporan_menunggu": laporan_menunggu,
        "dikembalikan": dikembalikan,
        "proses": hitung.get(S.PROSES, 0),
        "proses_bpsdm": hitung.get(S.PROSES_BPSDM, 0),
        "proses_pakln": hitung.get(S.PROSES_PAKLN, 0),
        "selesai": hitung.get(S.SELESAI, 0),
        "dibatalkan": hitung.get(S.DIBATALKAN, 0),
    }


def datatable(request, qs, order_fields, row_fn, search_q=None, default_order="-created_at"):
    """Respons JSON server-side DataTables (pola `users_data`)."""
    records_total = qs.count()
    search_value = request.GET.get("search[value]", "").strip()
    if search_value and search_q:
        qs = qs.filter(search_q(search_value)).distinct()
    records_filtered = qs.count()

    order_col = request.GET.get("order[0][column]")
    order_field = order_fields.get(order_col, default_order)
    if request.GET.get("order[0][dir]") == "desc" and not order_field.startswith("-"):
        order_field = f"-{order_field}"
    qs = qs.order_by(order_field, "-created_at")

    try:
        start = int(request.GET.get("start", 0))
        length = int(request.GET.get("length", 10))
    except ValueError:
        start, length = 0, 10
    page = qs[start:] if length == -1 else qs[start:start + length]
    return JsonResponse({
        "draw": int(request.GET.get("draw", 1)),
        "recordsTotal": records_total,
        "recordsFiltered": records_filtered,
        "data": [row_fn(p) for p in page],
    })


def search_pengajuan(nilai):
    return (
        Q(kode__icontains=nilai)
        | Q(pegawai__profile__nama__icontains=nilai)
        | Q(pegawai__profile__nip__icontains=nilai)
        | Q(pegawai__profile__unit_kerja__icontains=nilai)
        | Q(kategori__nama_kategori__icontains=nilai)
        | Q(tujuan_negara__nama_negara__icontains=nilai)
    )


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def status_pelaporan(p):
    if not p.is_pdln or p.status != S.SELESAI:
        return ""
    laporan = p.laporan
    if laporan:
        return laporan.get_status_display()
    if p.tgl_kembali and (timezone.localdate() - p.tgl_kembali).days > BATAS_HARI_LAPORAN:
        return "Belum Melapor (Terlambat)"
    return "Belum Melapor"


KOLOM_EXPORT = [
    "Kode", "Jenis", "Tipe", "Nama", "NIP", "Jabatan", "Unit Organisasi", "Kategori", "Negara",
    "Kota", "Penyelenggara/Perguruan Tinggi", "Beasiswa", "Sumber Pembiayaan", "Tgl Berangkat",
    "Tgl Kembali", "Mulai Kegiatan", "Selesai Kegiatan", "Jumlah Hari Kerja", "Maksud Perjalanan",
    "Kanal", "Status", "Tgl Pengajuan", "Tgl Masuk BPSDM", "Tgl Masuk PAKLN", "Tgl Selesai",
    "Perlu Visa", "Jumlah Dikembalikan", "Status Pelaporan", "Tgl Dibatalkan", "Pembatalan Diajukan Oleh",
    "Alasan Pembatalan",
]

_AKSI_KEMBALI = [Aksi.DIKEMBALIKAN_UNOR, Aksi.DIKEMBALIKAN_BPSDM, Aksi.DIKEMBALIKAN_PAKLN]


def baris_export(qs):
    qs = qs.select_related(
        "pegawai__profile", "unit_organisasi", "kategori", "sumber_pembiayaan", "detail_pdln", "laporan_pdln",
    ).prefetch_related("tujuan_negara").order_by("-tgl_pengajuan", "-created_at")
    ids = list(qs.values_list("id", flat=True))
    kembali = dict(
        RiwayatPengajuan.objects.filter(pengajuan_id__in=ids, aksi__in=_AKSI_KEMBALI)
        .values_list("pengajuan_id").annotate(n=Count("id")).values_list("pengajuan_id", "n")
    )
    batal = {
        m.pengajuan_id: m for m in PermohonanPembatalan.objects.filter(
            pengajuan_id__in=ids, status=PermohonanPembatalan.Status.DISETUJUI
        )
    }
    peran = dict(User.Role.choices)
    for p in qs:
        prof = getattr(p.pegawai, "profile", None)
        d = p.detail
        negara = list(p.tujuan_negara.all())
        m = batal.get(p.id)
        yield [
            p.kode, p.get_jenis_perjalanan_display(), p.get_tipe_pdln_display() if p.tipe_pdln else "",
            prof.nama if prof else p.pegawai.get_full_name(), prof.nip if prof else "", prof.jabatan if prof else "",
            (p.unit_organisasi.name if p.unit_organisasi else (prof.unit_organisasi.name if prof and prof.unit_organisasi else "")),
            p.kategori.nama_kategori if p.kategori else "", ", ".join(n.nama_negara for n in negara),
            d.kota_tujuan if d else "", (d.penyelenggara or d.perguruan_tinggi) if d else "",
            d.beasiswa_nama if d else "", p.sumber_pembiayaan.nama if p.sumber_pembiayaan else "",
            p.tgl_berangkat, p.tgl_kembali, d.tgl_mulai_kegiatan if d else None, d.tgl_selesai_kegiatan if d else None,
            p.jumlah_hari_kerja, p.maksud, p.get_kanal_display(), p.get_status_display(), p.tgl_pengajuan,
            p.tgl_masuk_bpsdm, p.tgl_masuk_pakln, p.tgl_selesai,
            ("Ya" if any(n.perlu_visa for n in negara) else "Tidak") if p.is_pdln else "",
            kembali.get(p.id, 0), status_pelaporan(p), p.tgl_dibatalkan,
            peran.get(m.diajukan_role, "") if m else "", m.alasan if m else "",
        ]


def _teks(v):
    if v is None:
        return ""
    if isinstance(v, date):
        return v.isoformat()
    return v


def respons_csv(qs, nama_berkas):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{nama_berkas}.csv"'
    response.write("﻿")  # BOM agar Excel membaca UTF-8
    writer = csv.writer(response)
    writer.writerow(KOLOM_EXPORT)
    for row in baris_export(qs):
        writer.writerow([_teks(v) for v in row])
    return response


def respons_xlsx(qs, nama_berkas, rekap_data=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(KOLOM_EXPORT)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in baris_export(qs):
        ws.append(row)
    for kolom in ws.columns:
        ws.column_dimensions[kolom[0].column_letter].width = 18

    if rekap_data:
        ws2 = wb.create_sheet("Ringkasan")
        for blok in rekap_data:
            ws2.append([blok["judul"]])
            ws2[ws2.max_row][0].font = Font(bold=True)
            ws2.append(blok["kolom"])
            for row in blok["baris"]:
                ws2.append(row)
            ws2.append([])

    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nama_berkas}.xlsx"'
    return response


# ---------------------------------------------------------------------------
# Rekap / pivot
# ---------------------------------------------------------------------------

_KOLOM_STATUS = [S.PROSES, S.PROSES_BPSDM, S.PROSES_PAKLN, S.SELESAI, S.DIBATALKAN, S.BELUM]


def _label_jenis(p_jenis, p_tipe):
    if p_jenis == JP.PDLN and p_tipe:
        return dict(T.choices).get(p_tipe, p_tipe)
    return "Non-Kedinasan"


def rekap(qs):
    """Daftar blok rekap {judul, kolom, baris}. Dipakai halaman Rekap dan
    sheet "Ringkasan" pada export XLSX."""
    label_status = dict(S.choices)
    blok = []

    # Per status
    data = defaultdict(Counter)
    for jenis, tipe, status in qs.values_list("jenis_perjalanan", "tipe_pdln", "status"):
        data[_label_jenis(jenis, tipe)][status] += 1
    kolom = ["Jenis / Tipe"] + [label_status[s] if s != S.BELUM else "Dikembalikan" for s in _KOLOM_STATUS] + ["Total", "% Selesai"]
    baris = []
    for nama, c in sorted(data.items()):
        total = sum(c.values())
        baris.append([nama] + [c[s] for s in _KOLOM_STATUS] + [total, f"{(c[S.SELESAI] * 100 // total) if total else 0}%"])
    blok.append({"kunci": "status", "judul": "Rekap per Status", "kolom": kolom, "baris": baris})

    # Per unit organisasi
    data = defaultdict(Counter)
    for unor, jenis, tipe in qs.values_list("unit_organisasi__name", "jenis_perjalanan", "tipe_pdln"):
        data[unor or "—"][_label_jenis(jenis, tipe)] += 1
    jenis_kolom = sorted({k for c in data.values() for k in c})
    blok.append({
        "kunci": "unor", "judul": "Rekap per Unit Organisasi",
        "kolom": ["Unit Organisasi"] + jenis_kolom + ["Total"],
        "baris": [[u] + [c[j] for j in jenis_kolom] + [sum(c.values())] for u, c in sorted(data.items())],
    })

    # Per negara
    data = defaultdict(Counter)
    for negara, jenis, tipe in qs.values_list("tujuan_negara__nama_negara", "jenis_perjalanan", "tipe_pdln"):
        data[negara or "—"][_label_jenis(jenis, tipe)] += 1
    jenis_kolom = sorted({k for c in data.values() for k in c})
    blok.append({
        "kunci": "negara", "judul": "Rekap per Negara Tujuan",
        "kolom": ["Negara"] + jenis_kolom + ["Total"],
        "baris": sorted(
            ([n] + [c[j] for j in jenis_kolom] + [sum(c.values())] for n, c in data.items()),
            key=lambda r: -r[-1],
        ),
    })

    # Per bulan (tgl pengajuan)
    data = defaultdict(Counter)
    for tgl, jenis, tipe in qs.values_list("tgl_pengajuan", "jenis_perjalanan", "tipe_pdln"):
        data[tgl.strftime("%Y-%m") if tgl else "—"][_label_jenis(jenis, tipe)] += 1
    jenis_kolom = sorted({k for c in data.values() for k in c})
    blok.append({
        "kunci": "bulan", "judul": "Rekap per Bulan Pengajuan",
        "kolom": ["Bulan"] + jenis_kolom + ["Total"],
        "baris": [[b] + [c[j] for j in jenis_kolom] + [sum(c.values())] for b, c in sorted(data.items())],
    })

    # Per kategori & sumber
    data = Counter()
    for kategori, sumber in qs.values_list("kategori__nama_kategori", "sumber_pembiayaan__nama"):
        data[(kategori or "—", sumber or "—")] += 1
    blok.append({
        "kunci": "kategori", "judul": "Rekap per Kategori & Sumber Pembiayaan",
        "kolom": ["Kategori", "Sumber Pembiayaan", "Jumlah"],
        "baris": [[k, s, n] for (k, s), n in sorted(data.items())],
    })

    blok.append(_rekap_waktu_proses(qs))
    blok.append(_rekap_pengembalian(qs))
    blok.append(_rekap_pembatalan(qs))
    return blok


def _rekap_waktu_proses(qs):
    """Lama (hari) tiap kunjungan ke sebuah tahap, dihitung dari riwayat:
    masuk = aksi dengan status_ke = tahap, keluar = aksi berikutnya dengan
    status_dari = tahap. Bolak-balik pengembalian ikut terhitung."""
    tahap_list = [S.PROSES, S.PROSES_BPSDM, S.PROSES_PAKLN]
    durasi = defaultdict(list)
    masuk = {}
    for pid, dari, ke, waktu in (
        RiwayatPengajuan.objects.filter(pengajuan__in=qs)
        .order_by("pengajuan_id", "created_at", "id")
        .values_list("pengajuan_id", "status_dari", "status_ke", "created_at")
    ):
        if dari != ke and (pid, dari) in masuk:
            durasi[dari].append((waktu - masuk.pop((pid, dari))).total_seconds() / 86400)
        if dari != ke and ke in tahap_list:
            masuk[(pid, ke)] = waktu
    label = dict(S.choices)
    return {
        "kunci": "waktu", "judul": "Waktu Proses per Tahap (hari kalender)",
        "kolom": ["Tahap", "Jumlah kunjungan", "Rata-rata", "Median", "Maksimum"],
        "baris": [
            [label[t], len(durasi[t]),
             round(sum(durasi[t]) / len(durasi[t]), 1) if durasi[t] else "—",
             round(median(durasi[t]), 1) if durasi[t] else "—",
             round(max(durasi[t]), 1) if durasi[t] else "—"]
            for t in tahap_list
        ],
    }


def _rekap_pengembalian(qs):
    total = qs.count()
    per_aksi = dict(
        RiwayatPengajuan.objects.filter(pengajuan__in=qs, aksi__in=_AKSI_KEMBALI)
        .values_list("aksi").annotate(n=Count("id")).values_list("aksi", "n")
    )
    pernah = qs.filter(riwayat__aksi__in=_AKSI_KEMBALI).distinct().count()
    label = dict(Aksi.choices)
    baris = [[label[a], per_aksi.get(a, 0)] for a in _AKSI_KEMBALI]
    baris.append(["Pengajuan yang pernah dikembalikan", f"{pernah} dari {total} ({(pernah * 100 // total) if total else 0}%)"])
    return {"kunci": "pengembalian", "judul": "Pengembalian", "kolom": ["Jenis Pengembalian", "Jumlah"], "baris": baris}


def _rekap_pembatalan(qs):
    data = defaultdict(Counter)
    for jenis, tipe, status in PermohonanPembatalan.objects.filter(pengajuan__in=qs).values_list(
        "pengajuan__jenis_perjalanan", "pengajuan__tipe_pdln", "status"
    ):
        data[_label_jenis(jenis, tipe)][status] += 1
    P = PermohonanPembatalan.Status
    kolom_status = [P.MENUNGGU_UNOR, P.MENUNGGU_PAKLN, P.DISETUJUI, P.DITOLAK, P.DITARIK]
    label = dict(P.choices)
    return {
        "kunci": "pembatalan", "judul": "Permohonan Pembatalan",
        "kolom": ["Jenis / Tipe"] + [label[s] for s in kolom_status] + ["Total"],
        "baris": [[j] + [c[s] for s in kolom_status] + [sum(c.values())] for j, c in sorted(data.items())],
    }
